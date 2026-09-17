"""Compact Batch 9 fixtures; no database, Redis, provider, or broker connections.

Run: python -B tests/test_upload_resource_limits.py
"""

import asyncio
from contextlib import contextmanager
from datetime import datetime, timezone
from dataclasses import replace
import importlib
import io
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch
import zipfile

import test_resource_admission as admission_harness


class UploadResourceTests(unittest.TestCase):
    __test__ = __name__ == "__main__"

    @classmethod
    def setUpClass(cls):
        admission_harness.ResourceAdmissionTests.setUpClass.__func__(cls)
        for attr, name in (("validation", "upload_validation"), ("ingress", "upload_ingress"),
                           ("resource_errors", "document_resource_errors"), ("chunks", "chunk_service"),
                           ("pdf", "pdf_service"), ("hybrid", "hybrid_pdf_service"), ("word", "word_service"),
                           ("excel", "excel_service"), ("txt", "text_service"), ("datalab", "datalab_service"),
                           ("advanced", "datalab_admission"), ("errors", "error_service"),
                           ("failures", "document_processing_errors")):
            setattr(cls, attr, importlib.import_module("app.services." + name))
        cls.worker = importlib.import_module("app.worker")

    def setUp(self):
        admission_harness.ResourceAdmissionTests.setUp(self)
        self.limits.upload_limits.cache_clear()
        self.folder = Path(self.stack.enter_context(tempfile.TemporaryDirectory(prefix="upload-resource-")))
        self.reject = self.resource_errors.DocumentResourceError

    def configure(self, **values):
        self.stack.enter_context(patch.dict(os.environ, {key: str(value) for key, value in values.items()}))
        self.limits.upload_limits.cache_clear()

    def source(self, name, data):
        path = self.folder / name
        path.write_bytes(data)
        return path

    def pdf_file(self, pages=1):
        from pypdf import PdfWriter
        path = self.folder / "synthetic.pdf"
        with PdfWriter() as writer:
            for _ in range(pages):
                writer.add_blank_page(width=72, height=72)
            writer.write(path)
        return path

    def office(self, extension=".docx", extra=None, body=None):
        stream = io.BytesIO()
        required = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr(required, body or "<document/>")
            for name, data in (extra or {}).items():
                entry = zipfile.ZipInfo(name)
                # ZipInfo normalizes backslashes on Windows; preserve hostile
                # on-disk names so the validator, not fixture creation, sees them.
                entry.filename = name
                entry.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(entry, data)
        stream.seek(0)
        return stream

    def assert_rejected(self, code, action):
        with self.assertRaises(self.reject) as error:
            action()
        self.assertEqual(error.exception.code, code)
        return error.exception

    def test_file_size_accepts_exact_limit_without_allocating_large_fixture(self):
        stream = MagicMock()
        stream.tell.return_value = 50 * 1024**2
        self.validation.validate_source_size(stream)
        stream.read.assert_not_called()

    def test_file_above_limit_rejected_before_reading(self):
        stream = MagicMock()
        stream.tell.return_value = 50 * 1024**2 + 1
        self.assert_rejected("file_size", lambda: self.validation.validate_source_size(stream))
        stream.read.assert_not_called()

    def test_empty_and_unsupported_uploads(self):
        self.assert_rejected("empty_file", lambda: self.validation.validate_document_source(io.BytesIO(), ".txt"))
        self.assert_rejected("unsupported_file", lambda: self.validation.validate_document_source(io.BytesIO(b"text"), ".exe"))

    def test_utf8_multibyte_across_old_sample_and_new_read_boundaries(self):
        for boundary in (8192, 65536):
            for bom in (b"", b"\xef\xbb\xbf"):
                text = "a" * (boundary - 1 - len(bom)) + "€ readable text"
                path = self.source("synthetic.txt", bom + text.encode())
                self.assertEqual(self.txt.extract_content_from_text(path)[0]["content"], text)

    def test_invalid_utf8_tail_and_nulls_rejected(self):
        for data in (b"good\xe2\x82", b"good\x00tail", b"good\xff"):
            self.assert_rejected("invalid_text", lambda: self.validation.validate_document_source(io.BytesIO(data), ".txt"))

    def test_pdf_below_and_above_page_limit(self):
        self.configure(MAX_PDF_PAGES=2)
        self.assertEqual(self.validation.validate_document_source(self.pdf_file(2), ".pdf"), 2)
        with patch.object(self.hybrid, "extract_content_with_datalab") as paid:
            self.assert_rejected("pdf_pages", lambda: self.hybrid.extract_content_from_hybrid_pdf(self.pdf_file(3)))
        paid.assert_not_called()

    def test_invalid_pdf_returns_safe_error(self):
        error = self.assert_rejected("invalid_pdf", lambda: self.validation.validate_document_source(io.BytesIO(b"%PDF-private-parser-secret"), ".pdf"))
        self.assertNotIn("secret", error.detail)
        self.assertEqual(error.status_code, 400)

    def test_standard_pdf_parser_retains_valid_behavior(self):
        self.assertEqual(self.pdf.extract_content_from_pdf(self.pdf_file()), [])

    def test_pdf_stream_bytes_rejected_before_text_parser(self):
        self.configure(MAX_PDF_PAGE_CONTENT_BYTES=4)
        stream = MagicMock()
        stream.get_object.return_value = stream
        stream.get_data.return_value = b"12345"
        page = MagicMock()
        page.get.return_value = stream
        self.assert_rejected("content_limit", lambda: self.validation.PDFTextBudget().extract(page))
        page.extract_text.assert_not_called()

    def test_datalab_batch_and_total_budget_are_distinct(self):
        self.configure(DATALAB_BATCH_SIZE=2, MAX_DATALAB_PAGES_PER_DOCUMENT=3)
        budget = self.advanced.AdvancedPageBudget("synthetic.pdf", [1, 2, 3])
        budget.consume("synthetic.pdf", "0-1")
        budget.consume("synthetic.pdf", "2")
        self.assertEqual(budget.used, {1, 2, 3})
        self.assert_rejected("advanced_pages", lambda: budget.consume("synthetic.pdf", "3"))
        self.assert_rejected("advanced_pages", lambda: self.advanced.AdvancedPageBudget("synthetic.pdf", [1, 2, 3, 4]))

    def test_datalab_rejects_oversized_duplicate_and_missing_admission(self):
        self.configure(DATALAB_BATCH_SIZE=2, MAX_DATALAB_PAGES_PER_DOCUMENT=3)
        budget = self.advanced.AdvancedPageBudget("synthetic.pdf", [1, 2, 3])
        self.assert_rejected("advanced_pages", lambda: budget.consume("synthetic.pdf", "0-2"))
        budget.consume("synthetic.pdf", "0")
        self.assert_rejected("advanced_pages", lambda: budget.consume("synthetic.pdf", "0"))
        with patch.object(self.datalab, "post_with_retry") as paid:
            self.assert_rejected("advanced_pages", lambda: self.datalab.convert_document_with_datalab("synthetic.pdf"))
        paid.assert_not_called()

    def test_scanned_overflow_fails_before_any_datalab_call(self):
        self.configure(MAX_DATALAB_PAGES_PER_DOCUMENT=1)
        with patch.object(self.hybrid, "extract_content_with_datalab") as paid:
            self.assert_rejected("advanced_pages", lambda: self.hybrid.extract_content_from_hybrid_pdf(self.pdf_file(2)))
        paid.assert_not_called()

    def test_hybrid_uses_only_admitted_pages_and_falls_back_deterministically(self):
        self.configure(DATALAB_BATCH_SIZE=2, MAX_DATALAB_PAGES_PER_DOCUMENT=3)
        ranges = []
        def response(data):
            result = MagicMock(ok=True)
            result.iter_content.return_value = [json.dumps(data).encode()]
            return result
        def submit(url, **kwargs):
            ranges.append(kwargs["data"]["page_range"])
            return response({"success": True, "request_check_url": "https://synthetic.invalid/check"})
        with patch.object(self.validation.PDFTextBudget, "extract", return_value="Readable standard text"), \
             patch.object(self.hybrid, "is_complex_page", return_value=True), \
             patch.object(self.datalab, "post_with_retry", side_effect=submit), \
             patch.object(self.datalab, "get_with_retry", side_effect=lambda *args: response({
                 "status": "complete", "json": {"children": [{"type": "text", "text": "Advanced text"}]}})), \
             patch.object(self.datalab.time, "sleep"), \
             patch.object(self.hybrid, "save_datalab_images") as images:
            blocks = self.hybrid.extract_content_from_hybrid_pdf(self.pdf_file(4))
        self.assertEqual(sorted(ranges), ["0-1", "2"])
        self.assertTrue(any(block["metadata"].get("page") == 4 and block["content"] == "Readable standard text" for block in blocks))
        images.assert_not_called()

    def test_datalab_response_size_is_bounded_and_closed(self):
        self.configure(MAX_DATALAB_RESPONSE_BYTES=4)
        response = MagicMock()
        response.iter_content.return_value = [b"123", b"45"]
        self.assert_rejected("content_limit", lambda: self.datalab.read_bounded_json(response))
        response.close.assert_called_once()

    def test_image_fallback_content_is_rejected_before_asset_uploads(self):
        path = self.pdf_file()
        self.configure(MAX_EXTRACTED_CHARACTERS=10)
        result = {"document_json": {"children": []}, "images": {"large-caption.png": "synthetic"}}
        with patch.object(self.hybrid, "classify_pdf_pages", return_value=([], [1])), \
             patch.object(self.hybrid, "extract_content_with_datalab", return_value=result), \
             patch.object(self.hybrid, "save_datalab_images") as save:
            self.assert_rejected("content_limit", lambda: self.hybrid.extract_content_from_hybrid_pdf(path))
        save.assert_not_called()

    def test_zip_entry_count_rejected_before_zipinfo_allocation(self):
        stream = self.office(extra={"extra.xml": "<extra/>"})
        self.configure(MAX_OFFICE_ZIP_ENTRIES=2)
        with patch.object(zipfile, "ZipFile") as archive:
            self.assert_rejected("office_expansion", lambda: self.validation.validate_document_source(stream, ".docx"))
        archive.assert_not_called()

    def test_forged_zip_entry_count_still_bounded(self):
        stream = self.office(extra={"extra.xml": "<extra/>"})
        data = bytearray(stream.getvalue())
        offset = data.rfind(b"PK\x05\x06")
        struct.pack_into("<2H", data, offset + 8, 1, 1)
        self.configure(MAX_OFFICE_ZIP_ENTRIES=2)
        with patch.object(zipfile, "ZipFile") as archive:
            self.assert_rejected("office_expansion", lambda: self.validation.validate_document_source(io.BytesIO(data), ".docx"))
        archive.assert_not_called()

    def test_zip_total_and_individual_expansion_limits(self):
        stream = self.office()
        self.configure(MAX_OFFICE_EXPANDED_BYTES=10)
        self.assert_rejected("office_expansion", lambda: self.validation.validate_document_source(stream, ".docx"))
        self.configure(MAX_OFFICE_EXPANDED_BYTES=100, MAX_OFFICE_ENTRY_BYTES=5)
        self.assert_rejected("office_expansion", lambda: self.validation.validate_document_source(stream, ".docx"))

    def test_zip_compression_ratio(self):
        stream = self.office(extra={"large.xml": "<x>" + "a" * 2000 + "</x>"})
        self.configure(MAX_OFFICE_COMPRESSION_RATIO=5)
        self.assert_rejected("office_expansion", lambda: self.validation.validate_document_source(stream, ".docx"))

    def test_zip_unsafe_paths_and_dtd_are_rejected(self):
        for name in ("../escape.xml", "/absolute.xml", "C:/drive.xml", "word\\escape.xml"):
            self.assert_rejected("invalid_office", lambda: self.validation.validate_document_source(self.office(extra={name: "<x/>"}), ".docx"))
        body = '<!DOCTYPE document [<!ENTITY x "test">]><document>&x;</document>'
        self.assert_rejected("invalid_office", lambda: self.validation.validate_document_source(self.office(body=body), ".docx"))

    def test_normal_docx_preserves_paragraph_content(self):
        from docx import Document
        document = Document()
        document.add_paragraph("A normal résumé €")
        path = self.folder / "normal.docx"
        document.save(path)
        self.assertEqual(self.word.extract_content_from_word(path)[0]["content"], "A normal résumé €")

    def test_docx_paragraph_and_table_scale(self):
        self.configure(MAX_DOCX_PARAGRAPHS=1)
        self.assert_rejected("docx_limit", lambda: self.validation.validate_document_source(self.office(body="<document><p/><p/></document>"), ".docx"))
        self.configure(MAX_DOCX_TABLES=1)
        self.assert_rejected("docx_limit", lambda: self.validation.validate_document_source(self.office(body="<document><tbl/><tbl/></document>"), ".docx"))

    def workbook(self, sheets=1):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active["A1"] = "Normal text"
        workbook.active["B1"] = "=1+2"
        for index in range(1, sheets):
            workbook.create_sheet(f"Sheet{index}")["A1"] = index
        path = self.folder / "normal.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_normal_xlsx_preserves_formulas_and_streaming_views(self):
        path = self.workbook()
        with patch.object(self.excel, "load_workbook", wraps=self.excel.load_workbook) as load:
            blocks = self.excel.extract_content_from_excel(path)
        self.assertEqual(blocks[0]["content"], "A1: Normal text | B1: Formula==1+2, Value=None")
        self.assertTrue(all(call.kwargs["read_only"] and not call.kwargs["keep_links"] for call in load.call_args_list))

    def test_xlsx_worksheet_limit_precedes_workbook_loading(self):
        path = self.workbook(2)
        self.configure(MAX_XLSX_WORKSHEETS=1)
        with patch.object(self.excel, "load_workbook") as load:
            self.assert_rejected("spreadsheet_limit", lambda: self.excel.extract_content_from_excel(path))
        load.assert_not_called()

    def test_xlsx_actual_rows_columns_and_grid_cells(self):
        cases = [({"MAX_XLSX_ROWS_PER_SHEET": 2}, '<row r="3"><c r="A3"/></row>'),
                 ({"MAX_XLSX_COLUMNS": 2}, '<row r="1"><c r="C1"/></row>'),
                 ({"MAX_XLSX_CELLS": 3}, '<row r="2"><c r="B2"/></row>')]
        for settings, row in cases:
            self.configure(**settings)
            stream = self.office(".xlsx", {"xl/worksheets/sheet1.xml": f'<worksheet><dimension ref="A1"/><sheetData>{row}</sheetData></worksheet>'})
            self.assert_rejected("spreadsheet_limit", lambda: self.validation.validate_document_source(stream, ".xlsx"))

    def test_xlsx_total_rows_across_sheets(self):
        self.configure(MAX_XLSX_TOTAL_ROWS=3)
        sheet = '<worksheet><sheetData><row r="2"><c r="A2"/></row></sheetData></worksheet>'
        stream = self.office(".xlsx", {f"xl/worksheets/sheet{i}.xml": sheet for i in (1, 2)})
        self.assert_rejected("spreadsheet_limit", lambda: self.validation.validate_document_source(stream, ".xlsx"))

    def test_extracted_character_utf8_byte_and_block_limits(self):
        self.configure(MAX_EXTRACTED_CHARACTERS=4)
        self.assert_rejected("content_limit", lambda: self.chunks.create_chunks_from_content([{"content": "12345"}]))
        self.configure(MAX_EXTRACTED_CHARACTERS=10, MAX_EXTRACTED_TEXT_BYTES=4)
        self.assert_rejected("content_limit", lambda: self.chunks.create_chunks_from_content([{"content": "€€"}]))
        self.configure(MAX_EXTRACTED_TEXT_BYTES=100, MAX_EXTRACTED_BLOCKS=1)
        self.assert_rejected("content_limit", lambda: self.chunks.create_chunks_from_content([{"content": "a"}, {"content": "b"}]))

    def test_chunk_budget_stops_incrementally_and_overlap_is_preserved(self):
        self.assertEqual(self.chunks.chunk_text("a b c d e", 3, 1), ["a b c", "c d e", "e"])
        self.configure(MAX_DOCUMENT_CHUNKS=2)
        budget = self.chunks.ContentBudget()
        self.assert_rejected("content_limit", lambda: self.chunks.chunk_text("a b c d e f g h", 2, 0, budget=budget))
        self.assertEqual(budget.chunks, 3)

    def test_resumed_checkpoint_is_checked_before_embedding(self):
        path = self.source("synthetic.txt", b"Normal")
        db = MagicMock()
        db.execute.return_value.one.return_value = (10_001, 5, 5)
        @contextmanager
        def session():
            yield db
        with patch.object(self.processing, "create_passage_embeddings") as paid:
            self.assert_rejected("content_limit", lambda: self.processing.process_admitted_document(
                SimpleNamespace(session=session), 1, str(path), "txt", True))
        paid.assert_not_called()

    def test_resource_failure_is_permanent_in_worker_and_releases_user_permit(self):
        path = self.source("synthetic.txt", b"Normal")
        document = SimpleNamespace(id=1, user_id=1, file_path=str(path), file_type="txt",
                                   processing_status="processing", processing_stage="uploaded")
        db = MagicMock()
        db.get.return_value = document
        db.scalar.return_value = None
        @contextmanager
        def session():
            yield db
        @contextmanager
        def claim(*args):
            connection = self.permits.connect()
            try:
                yield SimpleNamespace(session=session, connection=connection)
            finally:
                connection.close()
        self.configure(MAX_EXTRACTED_CHARACTERS=4)
        task = self.worker.process_document_task
        with patch.object(self.processing, "claim_document_processing", claim), \
             patch.object(self.processing, "inspect_embeddings", return_value=[SimpleNamespace(complete=False)]), \
             patch.object(self.processing, "extract_content", return_value=[{"content": "12345"}]), \
             patch.object(self.processing, "create_passage_embeddings") as paid, \
             patch.object(self.processing, "process_waiting_messages_for_document"), \
             patch.object(task, "retry") as retry, self.assertLogs(self.errors.logger, "ERROR"):
            self.assertEqual(task.run(1), "permanent_failure")
        self.assertEqual(document.processing_stage, "permanent_failure")
        self.assertEqual(document.processing_error, self.resource_errors.resource_messages()["content_limit"])
        self.assertEqual(self.permits.slots, {})
        paid.assert_not_called()
        retry.assert_not_called()
        db.add.assert_not_called()

    def test_permanent_resource_and_transient_retry_classification(self):
        for code in ("pdf_pages", "office_expansion", "spreadsheet_limit", "content_limit"):
            error = self.reject(code)
            error.__cause__ = TimeoutError("private provider secret")
            self.assertFalse(self.failures.is_retryable_processing_error(error))
        self.assertTrue(self.failures.is_retryable_processing_error(TimeoutError()))
        self.assertTrue(self.failures.is_retryable_processing_error(self.failures.RetryableDocumentProcessingError()))

    def http_client(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        app = FastAPI()
        app.include_router(self.documents.router)
        app.add_middleware(self.ingress.UploadBodyLimitMiddleware)
        app.add_exception_handler(self.reject, self.resource_errors.resource_validation_response)
        app.add_exception_handler(self.admission.ResourceRejected, self.admission.resource_error_response)
        self.db = MagicMock()
        self.db.get.return_value = SimpleNamespace(id=1)
        app.dependency_overrides[self.database.get_db] = lambda: self.db
        self.stack.enter_context(patch.object(self.auth, "decode_access_token", return_value={"sub": "1"}))
        self.quota_db = MagicMock()
        self.quota_db.execute.return_value.all.return_value = []
        def refresh(document):
            document.id = 1
            document.created_at = datetime.now(timezone.utc)
        self.quota_db.refresh.side_effect = refresh
        quota_session = self.stack.enter_context(patch.object(self.quota, "Session"))
        quota_session.return_value.__enter__.return_value = self.quota_db
        self.stack.enter_context(patch.object(self.documents, "UPLOAD_DIR", self.folder))
        self.dispatch = self.stack.enter_context(patch.object(self.documents, "dispatch_uploaded_document", side_effect=lambda **kwargs: kwargs["response"]))
        return self.stack.enter_context(TestClient(app, headers={"Authorization": "Bearer synthetic"}))

    def test_http_invalid_upload_releases_quota_and_does_not_save_or_dispatch(self):
        client = self.http_client()
        response = client.post("/documents", files={"file": ("invalid.pdf", b"private-parser-secret", "application/pdf")})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], "invalid_pdf")
        self.assertNotIn("secret", response.text)
        self.assertEqual(self.permits.slots, {})
        self.assertEqual(list(self.folder.iterdir()), [])
        self.quota_db.add.assert_not_called()
        self.quota_db.commit.assert_not_called()
        self.dispatch.assert_not_called()

    def test_http_valid_upload_saves_and_dispatches_once_then_releases(self):
        client = self.http_client()
        response = client.post("/documents", files={"file": ("normal.txt", b"Normal text", "text/plain")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["filename"], "normal.txt")
        self.assertEqual(len(list(self.folder.iterdir())), 1)
        self.quota_db.add.assert_called_once()
        self.quota_db.commit.assert_called_once()
        self.dispatch.assert_called_once()
        self.assertEqual(self.permits.slots, {})

    def test_write_time_size_rejection_removes_original_and_releases_quota(self):
        client = self.http_client()
        async def earlier_size_check(file):
            # Exercise the independent copy-time guard after content preflight.
            return None
        with patch.object(self.documents, "validate_file_size", side_effect=earlier_size_check), \
             patch.object(self.documents, "MAX_UPLOAD_SIZE_BYTES", 5):
            response = client.post("/documents", files={"file": ("normal.txt", b"123456", "text/plain")})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["code"], "file_size")
        self.assertEqual(list(self.folder.iterdir()), [])
        self.assertEqual(self.permits.slots, {})
        self.quota_db.add.assert_not_called()
        self.quota_db.commit.assert_not_called()
        self.dispatch.assert_not_called()

    def test_http_policy_exposes_only_configured_product_limits(self):
        self.configure(MAX_PDF_PAGES=12)
        response = self.http_client().get("/documents/upload-policy")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {"supported_extensions": [".pdf", ".docx", ".xlsx", ".txt"],
                                          "max_file_bytes": 50 * 1024**2, "max_pdf_pages": 12})

    def test_http_policy_cannot_be_shadowed_by_document_route_registered_first(self):
        from fastapi import APIRouter
        routes = sorted(self.documents.router.routes,
                        key=lambda route: route.endpoint is not self.documents.get_document)
        # Reproduce the unsafe ordering with the real routes and dependencies.
        with patch.object(self.documents, "router", APIRouter(routes=routes)), \
             patch.object(self.documents, "get_owned_document") as lookup:
            response = self.http_client().get("/documents/upload-policy")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), self.limits.upload_limits().public_policy())
        self.assertNotIn("int_parsing", response.text)
        lookup.assert_not_called()
        self.assertEqual(self.rates.calls, 1)

    def test_http_numeric_document_route_preserves_lookup_and_public_path(self):
        from fastapi import HTTPException
        client = self.http_client()
        document = SimpleNamespace(
            id=123, user_id=1, filename="synthetic.txt", file_type="txt", pages_count=None,
            processing_status="ready", processing_stage="completed", processing_progress=100,
            processing_error=None, created_at=datetime.now(timezone.utc),
        )
        with patch.object(self.documents, "get_owned_document", return_value=document) as lookup:
            response = client.get("/documents/123")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["id"], 123)
            lookup.assert_called_once_with(db=self.db, document_id=123, current_user=self.db.get.return_value)
            lookup.side_effect = HTTPException(404, "Document not found")
            missing = client.get("/documents/124")
        self.assertEqual(missing.status_code, 404)
        self.assertEqual(missing.json(), {"detail": "Document not found"})
        paths = client.get("/openapi.json").json()["paths"]
        self.assertIn("get", paths["/documents/{document_id}"])
        self.assertEqual(self.rates.calls, 2)

    def test_http_invalid_document_paths_fail_without_document_lookup(self):
        client = self.http_client()
        with patch.object(self.documents, "get_owned_document") as lookup:
            for value in ("not-a-number", "1.5", "-1"):
                with self.subTest(value=value):
                    response = client.get(f"/documents/{value}")
                    # The unchanged DELETE route matches the path only; no GET
                    # handler accepts it, so Starlette returns its safe 405.
                    self.assertEqual(response.status_code, 405, response.text)
                    self.assertEqual(response.json(), {"detail": "Method Not Allowed"})
        lookup.assert_not_called()

    def test_http_policy_and_document_routes_keep_auth_and_admission(self):
        client = self.http_client()
        with patch.object(self.documents, "get_owned_document") as lookup, \
             patch.object(self.documents, "upload_limits") as policy:
            for path in ("/documents/upload-policy", "/documents/123"):
                with self.subTest(path=path):
                    client.headers.pop("Authorization", None)
                    before = self.rates.calls
                    self.assertEqual(client.get(path).status_code, 401)
                    self.assertEqual(self.rates.calls, before)
                    client.headers["Authorization"] = "Bearer synthetic"
                    with patch.object(self.rates, "eval", return_value=(0, 7)):
                        limited = client.get(path)
                    self.assertEqual(limited.status_code, 429)
                    self.assertEqual(limited.json()["code"], "rate_limit")
                    self.assertEqual(limited.headers["Retry-After"], "7")
                    with patch.object(self.rates, "eval", side_effect=RuntimeError("private Redis marker")):
                        unavailable = client.get(path)
                    self.assertEqual(unavailable.status_code, 503)
                    self.assertEqual(unavailable.json()["code"], "admission_unavailable")
                    self.assertEqual(unavailable.headers["Retry-After"], "5")
                    self.assertNotIn("private Redis marker", unavailable.text)
        lookup.assert_not_called()
        policy.assert_not_called()

    def test_declared_oversized_body_rejected_before_multipart_or_quota(self):
        client = self.http_client()
        with patch.object(self.documents, "validate_file_content") as validation:
            response = client.post("/documents", content=b"tiny", headers={
                "Content-Length": str(51 * 1024**2 + 1), "Content-Type": "multipart/form-data; boundary=test"})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["code"], "request_size")
        self.quota_db.add.assert_not_called()
        validation.assert_not_called()

    def test_multipart_rejects_extra_files_and_fields(self):
        client = self.http_client()
        responses = [client.post("/documents", files=[("file", ("a.txt", b"a")), ("file", ("b.txt", b"b"))]),
                     client.post("/documents", files={"file": ("a.txt", b"a")}, data={"extra": "value"})]
        self.assertEqual([response.status_code for response in responses], [400, 400])
        self.quota_db.add.assert_not_called()
        self.assertEqual(self.permits.slots, {})

    def test_chunked_and_false_content_length_close_partial_spools(self):
        import starlette.formparsers as formparsers
        app = self.http_client().app
        settings = replace(self.limits.upload_limits(), file_bytes=128, multipart_overhead_bytes=64)
        real_spool = formparsers.SpooledTemporaryFile
        for length in (None, b"1"):
            with self.subTest(length=length):
                spools, sent = [], []
                chunks = iter([b'--test\r\nContent-Disposition: form-data; name="file"; filename="a.txt"\r\n\r\na', b"x" * 193])
                def spool(*args, **kwargs):
                    result = real_spool(*args, **kwargs)
                    spools.append(result)
                    return result
                async def receive():
                    return {"type": "http.request", "body": next(chunks), "more_body": True}
                async def send(message):
                    sent.append(message)
                headers = [(b"content-type", b"multipart/form-data; boundary=test"), (b"authorization", b"Bearer synthetic")]
                if length is not None:
                    headers.append((b"content-length", length))
                scope = {"type": "http", "http_version": "1.1", "asgi": {"version": "3.0"},
                         "method": "POST", "scheme": "http", "path": "/documents", "root_path": "",
                         "query_string": b"", "headers": headers, "client": ("127.0.0.1", 1), "server": ("test", 80)}
                with patch.object(self.ingress, "upload_limits", return_value=settings), patch.object(formparsers, "SpooledTemporaryFile", side_effect=spool):
                    asyncio.run(app(scope, receive, send))
                self.assertEqual(sent[0]["status"], 413)
                self.assertTrue(spools and all(file.closed for file in spools))
                self.assertEqual(self.permits.slots, {})


def test_upload_resources_in_isolated_process():
    result = subprocess.run([sys.executable, "-B", str(Path(__file__).resolve())],
                            cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, timeout=90)
    assert result.returncode == 0, result.stdout + result.stderr


if __name__ == "__main__":
    unittest.main(verbosity=2)
